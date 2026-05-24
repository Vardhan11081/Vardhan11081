from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx.opc.constants

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2)
    section.right_margin  = Cm(2)

# ── Default body font ──
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

def set_heading(para, text, level=1):
    para.text = text
    para.style = doc.styles[f'Heading {level}']
    run = para.runs[0]
    run.font.name = 'Calibri'
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    else:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2C, 0x5F, 0x8A)

def add_para(doc, text, bold_first=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = Pt(16.5)
    if bold_first:
        run = p.add_run(bold_first + " ")
        run.font.bold = True
        run.font.size = Pt(11)
    run2 = p.add_run(text)
    run2.font.size = Pt(11)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = Pt(16)
    if bold_prefix:
        r1 = p.add_run(bold_prefix + ": ")
        r1.font.bold = True
        r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    return p

# ──────────────────────────────────────────
# COVER / TITLE
# ──────────────────────────────────────────
title_p = doc.add_heading('', 0)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('BN314 – System Architecture\nAssignment 2: Case Study Analysis')
run.font.name  = 'Calibri'
run.font.size  = Pt(16)
run.font.bold  = True
run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_r = sub_p.add_run('FreshBite Salads – Online Ordering System\nQuestions 5 & 6')
sub_r.font.size  = Pt(12)
sub_r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

doc.add_paragraph()

# ──────────────────────────────────────────
# QUESTION 5 – First-Level DFD
# ──────────────────────────────────────────
h1 = doc.add_heading('', 1)
set_heading(h1, 'Question 5: First-Level Data Flow Diagram (Level-1 DFD)', 1)

add_para(doc, (
    "A Data Flow Diagram (DFD) is a graphical representation that models how data moves "
    "through a system. The Level-1 DFD expands the context diagram by decomposing the "
    "central system process into its major sub-processes, showing the data stores each "
    "process reads from or writes to, and the data flows that connect them."
))

h2 = doc.add_heading('', 2)
set_heading(h2, '5.1  Major Processes', 2)

processes = [
    ("P1 – User Registration & Authentication",
     "Handles new-user sign-up (collects username, name, DOB, email, password), "
     "enforces strong password rules, generates a verification code sent via email, "
     "and validates login credentials. After three failed attempts, it suspends the account."),
    ("P2 – Browse Salads",
     "Retrieves salad data (name, description, ingredients, nutritional info, allergen "
     "warnings, portion sizes, price) from the Salad Catalog and presents it to the customer."),
    ("P3 – Shopping Cart Management",
     "Allows customers to add, update, or remove items in a session-scoped cart. "
     "Persists cart contents to the Shopping Cart data store."),
    ("P4 – Checkout & Order Processing",
     "Calculates the order total including tax and delivery charges, reads the delivery "
     "address, creates a pending order record, and coordinates with Payment Processing. "
     "On successful payment it finalises the order and triggers confirmation emails."),
    ("P5 – Payment Processing",
     "Sends payment requests to the external Payment Gateway, receives the authorisation "
     "response, and returns a payment status to Checkout."),
    ("P6 – Order Tracking",
     "Updates and exposes the order lifecycle: Preparing → Out for Delivery → Delivered. "
     "Receives delivery status from the Delivery Service and notifies the customer."),
    ("P7 – Reward Points Management",
     "Awards 1 point per completed purchase, checks if the threshold of 15 points is "
     "reached to unlock a free salad, and updates the Reward Points database."),
]
for name, desc in processes:
    add_bullet(doc, desc, bold_prefix=name)

h2b = doc.add_heading('', 2)
set_heading(h2b, '5.2  Data Stores', 2)

stores = [
    ("D1 – Salad Catalog", "Stores all salad records (ID, name, ingredients, price, sizes)."),
    ("D2 – User Database", "Stores registered user profiles, credentials, and account status."),
    ("D3 – Shopping Cart", "Persists session-level cart items (salad ID, quantity, portion size)."),
    ("D4 – Order Database", "Records all orders with status, cost breakdown, and timestamps."),
    ("D5 – Address Database", "Stores user-saved billing and delivery addresses."),
    ("D6 – Reward Points DB", "Tracks accumulated and redeemed loyalty points per user."),
]
for name, desc in stores:
    add_bullet(doc, desc, bold_prefix=name)

h2c = doc.add_heading('', 2)
set_heading(h2c, '5.3  Diagram', 2)

doc.add_picture('/workspace/dfd_level1.png', width=Inches(6.5))
last_para = doc.paragraphs[-1]
last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph('Figure 1: Level-1 Data Flow Diagram – FreshBite Salads Online Ordering System')
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.runs[0].font.size = Pt(9)
cap.runs[0].font.italic = True

doc.add_paragraph()

# ──────────────────────────────────────────
# QUESTION 6 – Class Diagram
# ──────────────────────────────────────────
h1b = doc.add_heading('', 1)
set_heading(h1b, 'Question 6: UML Class Diagram', 1)

add_para(doc, (
    "The class diagram below models the static structure of the FreshBite Salads ordering "
    "system using UML notation. Each class is presented with its attributes (visibility, "
    "name, and type), its methods, and its associations to other classes, including "
    "multiplicity labels."
))

h2d = doc.add_heading('', 2)
set_heading(h2d, '6.1  Key Classes, Attributes & Methods', 2)

classes = [
    ("User",
     "Represents a registered customer.",
     ["-userId: String", "-username: String", "-firstName / lastName: String",
      "-dateOfBirth: Date", "-email: String", "-password: String",
      "-accountStatus: String  {Active | Suspended | Banned}",
      "-rewardPoints: int", "-loginAttempts: int"],
     ["+register()", "+login(): boolean", "+logout()", "+updateProfile()",
      "+reactivateAccount(): boolean"]),
    ("Salad",
     "Represents a menu item.",
     ["-saladId: String", "-name: String", "-description: String",
      "-ingredients: List<String>", "-nutritionalInfo: String",
      "-allergenWarnings: String", "-price: double", "-portionSizes: List<String>"],
     ["+getDetails(): String", "+getPrice(size): double", "+getSizes(): List<String>"]),
    ("CustomSalad  (extends Salad)",
     "Inherits from Salad; lets customers build their own salad.",
     ["-customIngredients: List<String>", "-customName: String"],
     ["+createCustomSalad()", "+addIngredient(i: String)", "+removeIngredient(i: String)"]),
    ("ShoppingCart",
     "Session-scoped container for selected items.",
     ["-cartId: String", "-sessionId: String", "-createdAt: DateTime", "-status: String"],
     ["+addItem(item: CartItem)", "+removeItem(itemId: String)",
      "+updateQuantity(id, qty)", "+getTotal(): double", "+clearCart()"]),
    ("CartItem",
     "A single line in the cart.",
     ["-itemId: String", "-quantity: int", "-portionSize: String", "-unitPrice: double"],
     ["+calculateSubtotal(): double", "+updateQuantity(q: int)"]),
    ("Order",
     "A confirmed purchase record.",
     ["-orderId: String", "-totalCost: double", "-tax: double",
      "-deliveryCharge: double", "-status: String", "-createdAt: DateTime"],
     ["+placeOrder()", "+updateStatus(s: String)", "+cancelOrder()",
      "+getOrderDetails(): String"]),
    ("Payment",
     "Tracks payment transaction details.",
     ["-paymentId: String", "-amount: double", "-method: String  {CreditCard|Debit|PayPal}",
      "-status: String", "-transactionId: String", "-paidAt: DateTime"],
     ["+processPayment(): boolean", "+verifyPayment(): boolean", "+refund()"]),
    ("Address",
     "A saved delivery or billing address.",
     ["-addressId: String", "-street/city/state/postCode/country: String",
      "-isDefault: boolean"],
     ["+addAddress()", "+updateAddress()", "+deleteAddress()"]),
    ("RewardPoints",
     "Tracks loyalty points for a user.",
     ["-pointsId: String", "-totalPoints: int",
      "-redeemedPoints: int", "-lastUpdated: DateTime"],
     ["+earnPoints(n: int)", "+redeemPoints(): boolean", "+getBalance(): int"]),
    ("Notification",
     "Email or system message dispatched to users.",
     ["-notifId: String", "-type: String", "-message: String", "-sentAt: DateTime"],
     ["+sendEmail()", "+sendVerificationCode()"]),
]

for cname, cdesc, attrs, meths in classes:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(cname)
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x2C, 0x5F, 0x8A)
    add_para(doc, cdesc, space_after=2)
    # attrs
    for a in attrs:
        ap = doc.add_paragraph(style='List Bullet')
        ap.paragraph_format.space_after = Pt(1)
        rr = ap.add_run(a)
        rr.font.size = Pt(10)
        rr.font.name = 'Courier New'
    # methods
    for m in meths:
        mp = doc.add_paragraph(style='List Bullet')
        mp.paragraph_format.space_after = Pt(1)
        rm = mp.add_run(m)
        rm.font.size = Pt(10)
        rm.font.name = 'Courier New'
        rm.font.color.rgb = RGBColor(0x00, 0x6B, 0x00)
    doc.add_paragraph()

h2e = doc.add_heading('', 2)
set_heading(h2e, '6.2  Associations & Multiplicity', 2)

assocs = [
    ("User → Address", "One-to-many (1..0..*)", "A user can save multiple delivery/billing addresses."),
    ("User → ShoppingCart", "One-to-one (1..1)", "Each user owns one active cart per session."),
    ("User → Order", "One-to-many (1..0..*)", "A user can place many orders over time."),
    ("User → RewardPoints", "One-to-one (1..1)", "Each user has exactly one rewards account."),
    ("User → Notification", "One-to-many (1..0..*)", "Users receive multiple system notifications."),
    ("ShoppingCart → CartItem", "One-to-many (1..1..*)", "A cart holds one or more line items."),
    ("CartItem → Salad", "Many-to-one (0..*.1)", "Many cart items can reference the same salad."),
    ("CustomSalad --|> Salad", "Inheritance (generalisation)", "CustomSalad extends Salad with custom ingredients."),
    ("Order → CartItem", "One-to-many (1..1..*)", "An order is composed of one or more cart items."),
    ("Order → Payment", "One-to-one (1..1)", "Each order is settled by exactly one payment record."),
    ("Order → Address", "Many-to-one (0..*.1)", "Each order links to one delivery address."),
    ("Order → Notification", "One-to-many (1..0..*)", "Order events (confirmed, dispatched) trigger notifications."),
]
for aname, mult, desc in assocs:
    add_bullet(doc, f"{mult} – {desc}", bold_prefix=aname)

h2f = doc.add_heading('', 2)
set_heading(h2f, '6.3  Diagram', 2)

doc.add_picture('/workspace/class_diagram.png', width=Inches(6.5))
last_para2 = doc.paragraphs[-1]
last_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap2 = doc.add_paragraph('Figure 2: UML Class Diagram – FreshBite Salads Online Ordering System')
cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap2.runs[0].font.size = Pt(9)
cap2.runs[0].font.italic = True

# ──────────────────────────────────────────
# REFERENCES
# ──────────────────────────────────────────
doc.add_paragraph()
h1c = doc.add_heading('', 1)
set_heading(h1c, 'References', 1)

refs = [
    "[1] I. Sommerville, Software Engineering, 10th ed. Pearson, 2016.",
    "[2] G. Booch, J. Rumbaugh, and I. Jacobson, The Unified Modeling Language User Guide, "
    "2nd ed. Addison-Wesley, 2005.",
    "[3] E. Yourdon, Modern Structured Analysis. Prentice-Hall, 1989.",
    "[4] Object Management Group (OMG), 'OMG Unified Modeling Language Specification,' "
    "v2.5.1, 2017. [Online]. Available: https://www.omg.org/spec/UML/2.5.1/",
]
for r in refs:
    rp = doc.add_paragraph(style='List Number')
    rp.paragraph_format.space_after = Pt(4)
    rr = rp.add_run(r)
    rr.font.size = Pt(10)

doc.save('/workspace/BN314_A2_Q5_Q6_Answer.docx')
print("Document saved: BN314_A2_Q5_Q6_Answer.docx")

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import io

# ═══════════════════════════════════════════════════════════════
# 1. Build the FDD diagram using matplotlib and save as PNG
# ═══════════════════════════════════════════════════════════════

def draw_box(ax, x, y, w, h, text, fontsize=8, color='#2C6FAC', text_color='white', radius=0.3):
    box = mpatches.FancyBboxPatch((x - w/2, y - h/2), w, h,
                                   boxstyle=f"round,pad={radius*0.1}",
                                   linewidth=1.2, edgecolor='#1a4a7a',
                                   facecolor=color)
    ax.add_patch(box)
    ax.text(x, y, text, ha='center', va='center', fontsize=fontsize,
            color=text_color, fontweight='bold', wrap=True,
            multialignment='center')

def draw_line(ax, x1, y1, x2, y2):
    ax.plot([x1, x2], [y1, y2], 'k-', linewidth=1.2)

fig, ax = plt.subplots(figsize=(14, 9))
ax.set_xlim(0, 14)
ax.set_ylim(0, 9)
ax.axis('off')
fig.patch.set_facecolor('#f8f9fa')

# ── Level 0: Root ──────────────────────────────────────────────
draw_box(ax, 7, 8.2, 3.2, 0.9, 'School\nManagement', fontsize=10, color='#1a4a7a')

# ── Level 1: 5 top-level functions ────────────────────────────
L1 = [
    (1.4,  6.5, 'Student\nServices'),
    (3.8,  6.5, 'Academic\nAffairs'),
    (7.0,  6.5, 'Finance &\nAccounting'),
    (10.2, 6.5, 'IT &\nInfrastructure'),
    (12.8, 6.5, 'Human\nResources'),
]
for x, y, lbl in L1:
    draw_box(ax, x, y, 2.2, 0.85, lbl, fontsize=8, color='#2C6FAC')
    draw_line(ax, x, y + 0.425, 7, 7.75)

# ── Level 2 under "Student Services" ──────────────────────────
L2_ss = [
    (0.5, 4.8, 'Enrolment\n& Registration'),
    (1.9, 4.8, 'Student\nSupport'),
]
for x, y, lbl in L2_ss:
    draw_box(ax, x, y, 1.6, 0.8, lbl, fontsize=7.5, color='#4a90d9')
    draw_line(ax, x, y + 0.4, 1.4, 6.07)

# ── Level 2 under "Academic Affairs" ──────────────────────────
L2_aa = [
    (3.0, 4.8, 'Course\nManagement'),
    (4.6, 4.8, 'Assessment\n& Grading'),
]
for x, y, lbl in L2_aa:
    draw_box(ax, x, y, 1.6, 0.8, lbl, fontsize=7.5, color='#4a90d9')
    draw_line(ax, x, y + 0.4, 3.8, 6.07)

# ── Level 2 under "Finance & Accounting" ──────────────────────
L2_fi = [
    (6.1, 4.8, 'Fees &\nPayments'),
    (7.9, 4.8, 'Budgeting\n& Reporting'),
]
for x, y, lbl in L2_fi:
    draw_box(ax, x, y, 1.6, 0.8, lbl, fontsize=7.5, color='#4a90d9')
    draw_line(ax, x, y + 0.4, 7.0, 6.07)

# ── Level 2 under "IT & Infrastructure" ───────────────────────
L2_it = [
    (9.4, 4.8, 'Network &\nSystems'),
    (11.0,4.8, 'LMS &\nPortals'),
]
for x, y, lbl in L2_it:
    draw_box(ax, x, y, 1.6, 0.8, lbl, fontsize=7.5, color='#4a90d9')
    draw_line(ax, x, y + 0.4, 10.2, 6.07)

# ── Level 2 under "Human Resources" ───────────────────────────
L2_hr = [
    (12.1,4.8, 'Staff\nRecruitment'),
    (13.5,4.8, 'Payroll &\nBenefits'),
]
for x, y, lbl in L2_hr:
    draw_box(ax, x, y, 1.6, 0.8, lbl, fontsize=7.5, color='#4a90d9')
    draw_line(ax, x, y + 0.4, 12.8, 6.07)

# ── Level 3 under "Enrolment & Registration" ──────────────────
L3_enrol = [
    (0.1, 3.1, 'Online\nApplication'),
    (1.0, 3.1, 'Subject\nEnrolment'),
]
for x, y, lbl in L3_enrol:
    draw_box(ax, x, y, 1.0, 0.75, lbl, fontsize=6.8, color='#7db8e8', text_color='#1a1a1a')
    draw_line(ax, x, y + 0.375, 0.5, 4.4)

# ── Level 3 under "Course Management" ─────────────────────────
L3_course = [
    (2.6, 3.1, 'Timetable\nScheduling'),
    (3.6, 3.1, 'Curriculum\nDesign'),
    (4.6, 3.1, 'Faculty\nAssignment'),
]
for x, y, lbl in L3_course:
    draw_box(ax, x, y, 1.1, 0.75, lbl, fontsize=6.5, color='#7db8e8', text_color='#1a1a1a')
    draw_line(ax, x, y + 0.375, 3.0, 4.4)

# Title
ax.text(7, 8.8, 'School Management – Functional Decomposition Diagram',
        ha='center', va='center', fontsize=11, fontweight='bold', color='#1a1a1a')

# Legend
legend_handles = [
    mpatches.Patch(color='#1a4a7a', label='Level 0 – Root'),
    mpatches.Patch(color='#2C6FAC', label='Level 1 – Major Functions'),
    mpatches.Patch(color='#4a90d9', label='Level 2 – Sub-Functions'),
    mpatches.Patch(color='#7db8e8', label='Level 3 – Processes'),
]
ax.legend(handles=legend_handles, loc='lower left', fontsize=8, framealpha=0.9)

plt.tight_layout()
fdd_path = '/workspace/fdd_school.png'
plt.savefig(fdd_path, dpi=150, bbox_inches='tight', facecolor=fig.get_facecolor())
plt.close()
print(f'FDD diagram saved: {fdd_path}')


# ═══════════════════════════════════════════════════════════════
# 2. Build the Word document
# ═══════════════════════════════════════════════════════════════

doc = Document()

# ── Styles helpers ──────────────────────────────────────────────
def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def body(doc, text, size=11):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(size)
    return p

def bullet(doc, text, size=11):
    p = doc.add_paragraph(text, style='List Bullet')
    p.style.font.size = Pt(size)
    return p

def numbered(doc, text, size=11):
    p = doc.add_paragraph(text, style='List Number')
    p.style.font.size = Pt(size)
    return p

def space(doc):
    doc.add_paragraph()

# ── Title block ─────────────────────────────────────────────────
t = doc.add_heading('BN314 – System Architecture', 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
s = doc.add_heading('Lecture 3: Requirement Modelling II – Laboratory Activity', 1)
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
space(doc)

# ════════════════════════════════════════════════════════════════
# Q1 – Questionnaire
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 1: Student Registration Process Questionnaire', 2)
body(doc,
     'The following questionnaire is designed to gather student feedback on the school '
     'registration process. Guidelines applied: questions are concise, unbiased, cover '
     'both open and closed formats, and progress from general to specific.')

space(doc)
doc.add_heading('Melbourne Institute of Technology – Student Registration Feedback Survey', 3)
body(doc, 'Instructions: Please answer all questions honestly. Your responses are anonymous and will be used to improve our registration system.')
space(doc)

questions = [
    ('Section A – General Information', None),
    ('Q1.', 'What is your current level of study?\n   ☐ Undergraduate   ☐ Postgraduate   ☐ Diploma   ☐ Certificate'),
    ('Q2.', 'How many times have you completed the registration process at this institution?\n   ☐ Once   ☐ Twice   ☐ Three or more times'),
    ('Section B – Registration Experience', None),
    ('Q3.', 'On a scale of 1–5, how easy was the online registration process to navigate?\n   1 (Very Difficult) – 2 – 3 – 4 – 5 (Very Easy)'),
    ('Q4.', 'Were you able to find and enrol in your required subjects without difficulty?\n   ☐ Yes   ☐ No   ☐ Partially'),
    ('Q5.', 'How long did the entire registration process take you?\n   ☐ Less than 30 minutes   ☐ 30–60 minutes   ☐ Over 1 hour'),
    ('Q6.', 'Did you encounter any technical issues during registration (e.g., system errors, login problems)?\n   ☐ Yes (please describe below)   ☐ No\n   _______________________________________________'),
    ('Q7.', 'Were the registration instructions and guidelines clear and easy to follow?\n   ☐ Strongly Agree   ☐ Agree   ☐ Neutral   ☐ Disagree   ☐ Strongly Disagree'),
    ('Section C – Support & Assistance', None),
    ('Q8.', 'Did you require assistance from student services during registration?\n   ☐ Yes   ☐ No'),
    ('Q9.', 'If yes, how responsive and helpful was the support you received?\n   ☐ Very Helpful   ☐ Somewhat Helpful   ☐ Not Helpful'),
    ('Section D – Open Feedback', None),
    ('Q10.', 'What did you like most about the registration process?\n   _______________________________________________'),
    ('Q11.', 'What improvements would you suggest to make the process better?\n   _______________________________________________'),
    ('Q12.', 'Any additional comments?\n   _______________________________________________'),
]

for label, text in questions:
    if text is None:
        doc.add_heading(label, 4)
    else:
        p = doc.add_paragraph()
        run = p.add_run(f'{label} ')
        run.bold = True
        run.font.size = Pt(10)
        p.add_run(text).font.size = Pt(10)

space(doc)

# ════════════════════════════════════════════════════════════════
# Q2 – Fill-in Form
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 2: Simple Fill-in Registration Form (≥ 5 Fields)', 2)
body(doc,
     'The following is a simple fill-in form that could be used as a printed or digital '
     'student registration data-capture form (designed using Microsoft Word table style):')
space(doc)

doc.add_heading('Melbourne Institute of Technology – Student Registration Form', 3)
space(doc)

table = doc.add_table(rows=10, cols=2)
table.style = 'Table Grid'
fields = [
    ('Full Name',              '________________________________________________'),
    ('Student ID',             '________________________________________________'),
    ('Date of Birth',          '____ / ____ / ________'),
    ('Email Address',          '________________________________________________'),
    ('Phone Number',           '________________________________________________'),
    ('Program / Course',       '________________________________________________'),
    ('Year of Study',          '☐ Year 1   ☐ Year 2   ☐ Year 3   ☐ Year 4'),
    ('Study Mode',             '☐ Full-time   ☐ Part-time   ☐ Online'),
    ('Emergency Contact Name', '________________________________________________'),
    ('Signature & Date',       'Signature: ____________________   Date: __ / __ / ____'),
]
for i, (field, placeholder) in enumerate(fields):
    row = table.rows[i]
    row.cells[0].text = field
    row.cells[1].text = placeholder
    for cell in row.cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(10)

space(doc)

# ════════════════════════════════════════════════════════════════
# Q3 – FDD Diagram
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 3: Functional Decomposition Diagram – School Management', 2)
body(doc,
     'The diagram below shows a Functional Decomposition Diagram (FDD) for a school, '
     'replacing the library example from Figure 4-8. The top-level function is '
     '"School Management", which decomposes into five major functional areas, each '
     'further broken down into sub-functions and processes.')
space(doc)
doc.add_picture(fdd_path, width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
space(doc)
body(doc,
     'The five top-level functions are: Student Services, Academic Affairs, Finance & '
     'Accounting, IT & Infrastructure, and Human Resources. Each decomposes further — '
     'for example, Student Services → Enrolment & Registration → Online Application / '
     'Subject Enrolment; Academic Affairs → Course Management → Timetable Scheduling / '
     'Curriculum Design / Faculty Assignment.')
space(doc)

# ════════════════════════════════════════════════════════════════
# Q4 – IT Industry News Site
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 4: IT Industry News Website Review', 2)
body(doc,
     'Site visited: TechCrunch (https://techcrunch.com)')
space(doc)
body(doc, 'What I liked:')
bullet(doc, 'TechCrunch publishes timely, well-sourced articles on emerging technologies, '
            'startup ecosystems, AI/ML developments, and enterprise software — highly relevant '
            'to IT industry professionals and students.')
bullet(doc, 'The site is well-organised with clear category navigation (AI, Security, Cloud, '
            'Apps, etc.), making it easy to find topics of interest quickly.')
bullet(doc, 'Articles include expert opinions, industry data, and links to original research, '
            'providing depth beyond surface-level news.')
bullet(doc, 'The newsletter and podcast options allow users to consume content in multiple formats.')
space(doc)
body(doc, "What I didn't like:")
bullet(doc, 'The site is advertisement-heavy; pop-up ads and auto-play videos disrupt the '
            'reading experience, especially on mobile devices.')
bullet(doc, 'Some content has a US-centric focus, with limited coverage of Asia-Pacific IT '
            'industry developments — less relevant for Australian-based students and professionals.')
bullet(doc, 'Access to some premium analysis articles requires a paid subscription.')
space(doc)

# ════════════════════════════════════════════════════════════════
# Q5 – FDD in Requirements Modelling
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 5: How FDDs Are Used in Requirements Modelling', 2)
body(doc,
     'A Functional Decomposition Diagram (FDD) is a top-down hierarchical model that '
     'breaks a complex system or organisation into progressively smaller, more manageable '
     'functions. In requirements modelling, FDDs serve several important roles:')
space(doc)
bullet(doc, 'Scope Definition: An FDD establishes and communicates the boundaries of a '
            'system. By identifying all major functions at the top level, analysts can confirm '
            'with stakeholders exactly what is (and is not) within the project scope.')
bullet(doc, 'Requirements Organisation: FDDs provide a logical structure for organising '
            'requirements. Each node in the diagram corresponds to a function that will '
            'require detailed requirements. This prevents requirements from being overlooked.')
bullet(doc, 'Communication Tool: FDDs are easy for non-technical stakeholders (managers, '
            'users) to understand. They bridge the gap between business needs and technical '
            'specifications without requiring IT expertise to interpret.')
bullet(doc, 'Foundation for Other Models: FDDs are typically created before Data Flow '
            'Diagrams (DFDs) and process specifications. The lowest-level functions in an '
            'FDD become the processes in a DFD, ensuring consistency across models.')
bullet(doc, 'Gap Analysis: By reviewing an FDD with subject-matter experts, analysts can '
            'identify missing functions, redundant processes, and opportunities for '
            'consolidation or automation.')
bullet(doc, 'Modular Design: FDDs support a divide-and-conquer approach, allowing large '
            'projects to be broken into independent modules that can be designed, developed, '
            'and tested separately.')
space(doc)

# ════════════════════════════════════════════════════════════════
# Q6 – Business Process Modelling (BPM)
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 6: Business Process Modelling (BPM)', 2)
body(doc,
     'Business Process Modelling (BPM) is a technique used to represent, analyse, and '
     'improve the workflows and processes within an organisation. It creates visual or '
     'formal representations of how work flows through a business — including who performs '
     'each step, what data or objects are involved, and how decisions are made.')
space(doc)
body(doc, 'Common BPM notations and tools include:')
bullet(doc, 'Business Process Model and Notation (BPMN) – the international standard for '
            'graphical process modelling, using symbols for tasks, events, gateways, and flows.')
bullet(doc, 'Swimlane Diagrams – show processes across different roles or departments using '
            'horizontal or vertical lanes, making accountability visible.')
bullet(doc, 'UML Activity Diagrams – used to model workflows in object-oriented systems.')
space(doc)
body(doc, 'How BPM can be used:')
bullet(doc, 'Requirements Elicitation: BPM models help analysts understand current ("as-is") '
            'processes before designing improved ("to-be") processes, ensuring new systems '
            'align with real business operations.')
bullet(doc, 'Process Improvement: By visualising workflows, organisations can identify '
            'bottlenecks, redundancies, delays, and inefficiencies that can be eliminated '
            'or automated.')
bullet(doc, 'System Design: BPM models translate business workflows into system requirements, '
            'showing which processes need software support, what data must be captured, and '
            'how users interact with the system.')
bullet(doc, 'Communication: BPM diagrams serve as a shared language between business '
            'stakeholders and IT teams, reducing misunderstandings during development.')
bullet(doc, 'Compliance and Auditing: Documented business processes help organisations '
            'demonstrate regulatory compliance and provide audit trails.')
bullet(doc, 'Change Management: When implementing new systems, BPM models show employees '
            'exactly how their workflows will change, making training and adoption easier.')
space(doc)

# ════════════════════════════════════════════════════════════════
# Q7 – Data Flow Diagrams in Requirements Modelling
# ════════════════════════════════════════════════════════════════
heading(doc, 'Question 7: How Data Flow Diagrams (DFDs) Are Used in Requirements Modelling', 2)
body(doc,
     'A Data Flow Diagram (DFD) is a graphical representation that shows how data moves '
     'through an information system — from external entities (sources/sinks), through '
     'processes, to data stores, and back. DFDs use four components:')
space(doc)
bullet(doc, 'External Entities (squares/rectangles) – people or systems that send or '
            'receive data (e.g., Student, Lecturer, Admissions Office).')
bullet(doc, 'Processes (circles/rounded rectangles) – functions that transform input data '
            'into output data (e.g., "Validate Enrolment", "Generate Timetable").')
bullet(doc, 'Data Stores (open-ended rectangles) – repositories where data is held '
            '(e.g., Student Database, Course Register).')
bullet(doc, 'Data Flows (arrows) – the movement of data between entities, processes, '
            'and data stores, labelled with the data being transferred.')
space(doc)
body(doc, 'How DFDs are used in requirements modelling:')
bullet(doc, 'Context Diagrams (Level 0 DFD): A single-process diagram showing the entire '
            'system and its interactions with external entities. This defines the system '
            'boundary and is invaluable for agreeing scope with stakeholders early in the '
            'project.')
bullet(doc, 'Detailed Process Modelling (Level 1, 2 DFDs): Each process from the context '
            'diagram is exploded into more detailed sub-processes, making data requirements '
            'explicit. Analysts can specify exactly what data each process needs as input '
            'and produces as output.')
bullet(doc, 'Identifying Data Requirements: DFDs reveal what data must be stored, retrieved, '
            'and transformed — forming the basis for data dictionaries and database design.')
bullet(doc, 'Detecting Missing or Redundant Processes: DFDs highlight data flows that have '
            'no destination (data sinks with no process) or processes that receive no input '
            '— indicating incomplete or incorrect requirements.')
bullet(doc, 'Bridging Business and Technical Requirements: DFDs are accessible enough for '
            'business users to validate while being precise enough for developers to use '
            'as a design input.')
bullet(doc, 'Supporting Structured Analysis: DFDs form a central artefact in Structured '
            'Systems Analysis and Design (SSAD), working alongside FDDs, data dictionaries, '
            'and process specifications to produce a complete requirements specification.')
space(doc)

# ── Save ────────────────────────────────────────────────────────
doc.save('/workspace/lab3.docx')
print('lab3.docx created successfully.')

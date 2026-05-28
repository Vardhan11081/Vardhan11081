from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Title ──────────────────────────────────────────────────────────────────
title = doc.add_heading('BN314/MN611 – System Architecture', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_heading('Week 2 Laboratory Activity', level=1)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()

# Helper: add a numbered question heading
def q_heading(doc, number, text):
    p = doc.add_heading(f'Question {number}: {text}', level=2)
    return p

# Helper: add a body paragraph
def body(doc, text):
    p = doc.add_paragraph(text)
    p.style.font.size = Pt(11)
    return p

# Helper: add a bullet item
def bullet(doc, text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.style.font.size = Pt(11)
    return p

# ── Q1 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 1, 'Five Fact-Finding Questions & the Zachman Framework')
body(doc, 'The five classic fact-finding questions are:')
for item in ['Who? – identifies the people, roles, and stakeholders involved.',
             'What? – identifies the data, objects, or things being handled.',
             'Where? – identifies the locations where processes or data reside.',
             'When? – identifies timing, schedules, and events.',
             'How? – identifies the processes and procedures used.']:
    bullet(doc, item)
body(doc, 'The Zachman Framework adds a sixth question:')
bullet(doc, 'Why? – identifies the motivations, business rules, goals, and strategies behind the system.')
body(doc, 'Yes, "Why?" is important. It forces analysts to understand the business purpose and strategic intent behind a system, not just its mechanics. Without understanding why a system exists, solutions may be technically correct but fail to deliver real business value.')

doc.add_paragraph()

# ── Q2 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 2, 'Systems Requirements and Classification')
body(doc, 'A systems requirement is a characteristic or capability that a system must have to satisfy a business need, objective, or user expectation. Requirements define what a system must do or the constraints under which it must operate.')
body(doc, 'Systems requirements are classified into two main categories:')
bullet(doc, 'Functional Requirements – describe what the system must do; the specific behaviours, functions, and features (e.g., "the system shall generate monthly sales reports").')
bullet(doc, 'Non-Functional Requirements – describe how the system performs its functions; these include:')
for sub_item in ['  • Performance – response time, throughput, capacity.',
                 '  • Security – access control, data integrity.',
                 '  • Reliability – uptime, fault tolerance.',
                 '  • Usability – ease of use, accessibility.',
                 '  • Scalability – ability to handle growth.']:
    body(doc, sub_item)

doc.add_paragraph()

# ── Q3 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 3, 'JAD, RAD, and Team-Based Methods')
body(doc, 'Joint Application Development (JAD) is a structured, workshop-based approach where users, managers, and IT professionals collaborate intensively to define system requirements. Facilitated sessions replace multiple one-on-one interviews.')
body(doc, 'Rapid Application Development (RAD) is an iterative development methodology that uses prototyping, time-boxed phases, and active user participation to build systems quickly and incrementally.')
body(doc, 'How they differ from traditional fact-finding:')
bullet(doc, 'Traditional methods (interviews, questionnaires, observation) are sequential and analyst-driven; users are interviewed individually.')
bullet(doc, 'JAD/RAD bring all stakeholders together simultaneously, compressing the requirements phase and enabling real-time consensus.')
body(doc, 'Main advantages of team-based methods:')
bullet(doc, 'Faster requirements gathering through parallel discussion rather than sequential interviews.')
bullet(doc, 'Higher-quality requirements due to immediate clarification and cross-functional input.')
bullet(doc, 'Greater user ownership and buy-in, leading to smoother implementation.')
bullet(doc, 'Reduced misunderstandings and rework by resolving conflicts early.')

doc.add_paragraph()

# ── Q4 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 4, 'Total Cost of Ownership (TCO)')
body(doc, 'Total Cost of Ownership (TCO) is the comprehensive assessment of all direct and indirect costs associated with acquiring, implementing, operating, and eventually retiring an information system over its entire life cycle.')
body(doc, 'TCO includes:')
bullet(doc, 'Acquisition/development costs (hardware, software, licences).')
bullet(doc, 'Implementation costs (installation, configuration, data migration).')
bullet(doc, 'Operational costs (maintenance, support, utilities).')
bullet(doc, 'Training and personnel costs.')
bullet(doc, 'Downtime and opportunity costs.')
body(doc, 'Costs that are frequently underestimated:')
bullet(doc, 'End-user training and productivity loss during transition.')
bullet(doc, 'Ongoing maintenance, patches, and upgrades.')
bullet(doc, 'Technical support and help-desk operations.')
bullet(doc, 'Data conversion and integration with legacy systems.')
bullet(doc, 'Hidden costs such as system downtime, workarounds, and business disruption.')

doc.add_paragraph()

# ── Q5 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 5, 'Examples of Question Types')

body(doc, 'Closed-Ended Questions (yes/no or single-choice answers):')
bullet(doc, 'Do you currently use the existing inventory system daily?')
bullet(doc, 'Is the current report generated automatically or manually?')
bullet(doc, 'Does the system currently allow role-based access control?')

body(doc, 'Open-Ended Questions (invite detailed, narrative responses):')
bullet(doc, 'How would you describe the biggest challenges you face with the current system?')
bullet(doc, 'What improvements would make your daily workflow more efficient?')
bullet(doc, 'Can you walk me through how you process a customer order from start to finish?')

body(doc, 'Range-of-Response Questions (scaled or ranked answers):')
bullet(doc, 'On a scale of 1–10, how satisfied are you with the system\'s response time?')
bullet(doc, 'How often do you encounter errors in the system? (Never / Rarely / Sometimes / Often / Always)')
bullet(doc, 'Rank the following features in order of importance to your work: speed, accuracy, ease of use, reporting.')

doc.add_paragraph()

# ── Q6 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 6, 'Three Types of Sampling')
body(doc, 'The three types of sampling used in systems analysis are:')
bullet(doc, 'Random Sampling – every member of the population has an equal chance of being selected. Provides unbiased results for large, homogeneous populations.')
bullet(doc, 'Systematic Sampling – every nth record or item is selected (e.g., every 10th transaction). Practical and easy to implement.')
bullet(doc, 'Stratified Sampling – the population is divided into subgroups (strata) based on a characteristic, and samples are drawn from each stratum proportionally.')
body(doc, 'For analysing data input errors, stratified sampling would be most appropriate. Input errors are unlikely to be evenly distributed across all data entry types, users, or time periods. By stratifying the data (e.g., by department, data type, or operator), the analyst can ensure that error-prone subgroups are adequately represented and patterns can be identified accurately.')

doc.add_paragraph()

# ── Q7 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 7, 'The Hawthorne Effect')
body(doc, 'The Hawthorne Effect refers to the phenomenon where individuals modify their behaviour simply because they know they are being observed. The name comes from productivity studies conducted at the Western Electric Hawthorne Works in the 1920s–1930s, where workers increased output regardless of environmental changes, apparently due to the attention they received.')
body(doc, 'In systems analysis, the Hawthorne Effect is important because observing users during their normal work tasks can produce inaccurate data — they may work more carefully, follow procedures they normally skip, or alter their pace when they know an analyst is watching. This can lead to requirements or process documentation that does not reflect actual day-to-day operations.')
body(doc, 'Personal experience: The Hawthorne Effect is commonly experienced in academic settings — for example, when a lecturer or examiner observes a lab session, students tend to follow instructions more precisely and ask fewer casual questions than they would otherwise. It is also evident in workplace performance reviews, where employees perform at a higher or different standard when they know they are being evaluated.')

doc.add_paragraph()

# ── Q8 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 8, 'Functional Decomposition Diagram (FDD)')
body(doc, 'A Functional Decomposition Diagram (FDD) is a hierarchical chart that breaks a complex system or business function into progressively smaller, more detailed sub-functions. It visually represents the structure of a system\'s functions from the highest level down to individual tasks.')
body(doc, 'Why use an FDD?')
bullet(doc, 'Provides a high-level overview and detailed breakdown of system functions simultaneously.')
bullet(doc, 'Helps stakeholders understand scope and identify missing or redundant functions.')
bullet(doc, 'Serves as a foundation for other models such as data flow diagrams (DFDs) and process specifications.')
bullet(doc, 'Facilitates modular design and task assignment during development.')
body(doc, 'How to create an FDD:')
bullet(doc, 'Step 1 – Identify the main system or business function and place it at the top of the diagram (Level 0).')
bullet(doc, 'Step 2 – Break the top-level function into its major sub-functions (Level 1). Each sub-function should represent a distinct area of functionality.')
bullet(doc, 'Step 3 – Decompose each Level 1 function further into more specific tasks or processes (Level 2), continuing downward until the lowest-level functions are atomic (cannot be broken down further meaningfully).')
bullet(doc, 'Step 4 – Connect each level with lines showing the parent-child relationship.')
bullet(doc, 'Step 5 – Review the diagram with stakeholders to verify completeness, accuracy, and correct scope.')

doc.add_paragraph()

# ── Q9 ─────────────────────────────────────────────────────────────────────
q_heading(doc, 9, 'Agile Methods')
body(doc, 'Agile methods are a group of iterative, incremental software development approaches that prioritise:')
bullet(doc, 'Working software over comprehensive documentation.')
bullet(doc, 'Customer collaboration over contract negotiation.')
bullet(doc, 'Responding to change over following a fixed plan.')
bullet(doc, 'Individuals and interactions over processes and tools.')
body(doc, 'Common agile frameworks include Scrum, Kanban, Extreme Programming (XP), and SAFe.')
body(doc, 'Are agile methods better than traditional methods?')
body(doc, 'Agile is not universally better — it depends on the project context:')
bullet(doc, 'Agile is better when: requirements are uncertain or likely to change, fast delivery of working software is critical, and users can be actively involved throughout development.')
bullet(doc, 'Traditional (waterfall) methods are better when: requirements are well-defined and stable upfront, regulatory compliance demands extensive documentation, or the project has fixed scope and budget (e.g., government contracts, safety-critical systems).')
body(doc, 'In practice, many organisations use hybrid approaches that combine agile flexibility with traditional planning rigour.')

doc.add_paragraph()

# ── Q10 ────────────────────────────────────────────────────────────────────
q_heading(doc, 10, 'Presentation Audiences')
body(doc, 'Three different audiences an analyst might present to:')

body(doc, '1. Management / Executives')
bullet(doc, 'Focus: strategic impact, ROI, cost, risk, and high-level timelines.')
bullet(doc, 'Style: concise, visually rich (charts, dashboards), business language — avoid technical jargon.')
bullet(doc, 'Goal: gain approval and funding.')

body(doc, '2. Technical Team / IT Staff')
bullet(doc, 'Focus: system architecture, data models, APIs, infrastructure, and implementation details.')
bullet(doc, 'Style: detailed, technical, use diagrams such as ERDs, DFDs, and UML. Open to deep discussion.')
bullet(doc, 'Goal: ensure shared understanding for implementation.')

body(doc, '3. End Users / Operational Staff')
bullet(doc, 'Focus: how the new system will affect their daily tasks, workflows, and responsibilities.')
bullet(doc, 'Style: simple language, demos, screenshots, step-by-step walkthroughs. Empathetic tone.')
bullet(doc, 'Goal: build confidence, address concerns, and encourage adoption.')

body(doc, 'How presentations differ:')
bullet(doc, 'For executives — what and why (business value).')
bullet(doc, 'For technical teams — how (implementation details).')
bullet(doc, 'For end users — what changes for me and how do I use it.')

body(doc, 'Most challenging: Presenting to end users is often the most challenging because they may have varying levels of technical literacy, strong emotional reactions to change, and concerns about job security. Managing resistance, simplifying complex concepts, and maintaining engagement requires both technical knowledge and strong interpersonal communication skills.')

# ── Save ───────────────────────────────────────────────────────────────────
doc.save('/workspace/lab2.docx')
print('lab2.docx created successfully.')

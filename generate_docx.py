"""Generate Hema_Vardhan_Moganti_Resume.docx from structured resume data."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Colours ──────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x0F, 0x34, 0x60)
MID_BLUE    = RGBColor(0x16, 0x57, 0x9A)
BODY_GREY   = RGBColor(0x33, 0x33, 0x44)
LIGHT_GREY  = RGBColor(0x66, 0x66, 0x66)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)

def set_cell_bg(cell, hex_color: str):
    """Fill a table cell background with a hex colour (e.g. '0F3460')."""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_para_border_bottom(para, color='0F3460', sz='6'):
    """Draw a bottom border under a paragraph (used for section headings)."""
    pPr   = para._p.get_or_add_pPr()
    pBdr  = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    sz)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def remove_table_borders(table):
    """Remove all cell borders from a table."""
    tbl  = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),   'none')
        el.set(qn('w:sz'),    '0')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
        tblBorders.append(el)
    tblPr.append(tblBorders)

def add_section_heading(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text.upper())
    run.bold      = True
    run.font.size = Pt(9)
    run.font.color.rgb = DARK_BLUE
    set_para_border_bottom(p)
    return p

def add_body(doc, text: str, bold=False, italic=False, size=10, color=BODY_GREY, indent=0, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before    = Pt(0)
    p.paragraph_format.space_after     = Pt(space_after)
    p.paragraph_format.left_indent     = Cm(indent)
    run = p.add_run(text)
    run.bold        = bold
    run.italic      = italic
    run.font.size   = Pt(size)
    run.font.color.rgb = color
    return p

def bullet(doc, text: str, level=0, size=10, bold_prefix: str = None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.left_indent  = Cm(0.5 + level * 0.4)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.size = Pt(size)
        r1.font.color.rgb = DARK_BLUE
        r2 = p.add_run(text)
        r2.font.size = Pt(size)
        r2.font.color.rgb = BODY_GREY
    else:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.color.rgb = BODY_GREY
    return p

# ════════════════════════════════════════════════════════════════
doc = Document()

# ── Page margins ─────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

# ════════════════════════════════════════════════════════════════
# HEADER  (dark-blue banner table)
# ════════════════════════════════════════════════════════════════
hdr_table = doc.add_table(rows=1, cols=2)
hdr_table.alignment = WD_TABLE_ALIGNMENT.LEFT
remove_table_borders(hdr_table)
hdr_table.columns[0].width = Inches(4.8)
hdr_table.columns[1].width = Inches(2.0)

left_cell  = hdr_table.cell(0, 0)
right_cell = hdr_table.cell(0, 1)
set_cell_bg(left_cell,  '0F3460')
set_cell_bg(right_cell, '0F3460')

for cell in (left_cell, right_cell):
    cell.paragraphs[0].clear()

# Name
name_p = left_cell.add_paragraph()
name_p.paragraph_format.space_before = Pt(8)
name_p.paragraph_format.space_after  = Pt(2)
name_run = name_p.add_run('Hema Vardhan Moganti')
name_run.bold           = True
name_run.font.size      = Pt(20)
name_run.font.color.rgb = WHITE

# Tagline
tag_p  = left_cell.add_paragraph()
tag_p.paragraph_format.space_before = Pt(0)
tag_p.paragraph_format.space_after  = Pt(8)
tag_run = tag_p.add_run('Cloud & AI Solutions Developer  |  Azure & Multi-Cloud  |  Cybersecurity')
tag_run.font.size      = Pt(10)
tag_run.font.color.rgb = RGBColor(0x7E, 0xB8, 0xF7)

# Contact lines
contacts = [
    ('📞 ', '+61 040 1257 633'),
    ('✉  ', 'Vardhanmoga11081@gmail.com'),
    ('📍 ', '101 Ballarat Rd, Footscray, VIC 3011'),
    ('🔗 ', 'linkedin.com/in/hema-vardhan-moganti'),
]
for icon, val in contacts:
    cp = left_cell.add_paragraph()
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after  = Pt(1)
    ic = cp.add_run(icon)
    ic.font.size = Pt(9)
    ic.font.color.rgb = RGBColor(0x7E, 0xB8, 0xF7)
    vr = cp.add_run(val)
    vr.font.size = Pt(9)
    vr.font.color.rgb = RGBColor(0xD0, 0xE8, 0xFF)

# Right cell – location
loc_p = right_cell.add_paragraph()
loc_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
loc_p.paragraph_format.space_before = Pt(20)
lr = loc_p.add_run('Melbourne, VIC, Australia\nUnrestricted Work Rights')
lr.font.size      = Pt(9)
lr.font.color.rgb = RGBColor(0xA8, 0xC8, 0xEE)

doc.add_paragraph().paragraph_format.space_after = Pt(2)  # small spacer

# ════════════════════════════════════════════════════════════════
# PROFESSIONAL SUMMARY
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Professional Summary')
add_body(doc,
    "Final-year Bachelor of Networking (Cyber Security) student with hands-on cloud and AI development experience "
    "spanning Microsoft Azure, AWS, and Google Cloud. Demonstrated ability to design, deploy, and manage cloud-native "
    "solutions — including an AI-powered multi-cloud text recognition and image enhancement system — and secure "
    "enterprise networks. Proficient in Python, Bash, and PowerShell scripting with practical exposure to CI/CD "
    "pipelines, Git workflows, Infrastructure as Code, and DevOps practices. Passionate about Azure PaaS services, "
    "proactive monitoring, automation, and delivering high-quality outcomes for customers in managed-services "
    "environments.",
    size=10)

# ════════════════════════════════════════════════════════════════
# TECHNICAL SKILLS
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Technical Skills')

skills = [
    ('Cloud Platforms:',       'Microsoft Azure, AWS, Google Cloud Platform'),
    ('Azure Services:',        'Azure VM, Azure Functions, App Service, Azure Monitor, Log Analytics, Active Directory'),
    ('AWS Services:',          'Lambda, S3, Rekognition, EventBridge, CloudWatch'),
    ('GCP Services:',          'Vision API, Cloud Functions, Cloud Storage, Cloud Run'),
    ('Programming:',           'Python, Bash, PowerShell, HTML, CSS, SQL'),
    ('DevOps & IaC:',          'Git, GitHub Actions, CI/CD pipelines, Bicep/ARM (foundational), Infrastructure as Code'),
    ('AI / ML:',               'Computer Vision (OCR/text detection), OpenCV, Pillow, cloud-native AI APIs'),
    ('Networking & Security:', 'TCP/IP, DNS, DHCP, VPN, VLANs, Active Directory, Group Policy, Wireshark'),
    ('Tools:',                 'Cisco Packet Tracer, Cisco Modelling Labs, Autopsy, FTK Imager, VS Code, PyCharm, SQLite, MS Access'),
    ('Operating Systems:',     'Windows Server 2022, Windows 10/11, Linux (Ubuntu)'),
]

for label, val in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    lr = p.add_run(label + '  ')
    lr.bold           = True
    lr.font.size      = Pt(10)
    lr.font.color.rgb = DARK_BLUE
    vr = p.add_run(val)
    vr.font.size      = Pt(10)
    vr.font.color.rgb = BODY_GREY

# ════════════════════════════════════════════════════════════════
# EDUCATION
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Education')
edu_tbl = doc.add_table(rows=1, cols=2)
remove_table_borders(edu_tbl)
edu_tbl.columns[0].width = Inches(4.5)
edu_tbl.columns[1].width = Inches(2.3)

lc = edu_tbl.cell(0,0)
rc = edu_tbl.cell(0,1)
lc.paragraphs[0].clear(); rc.paragraphs[0].clear()

p1 = lc.add_paragraph()
p1.paragraph_format.space_before = Pt(0); p1.paragraph_format.space_after = Pt(1)
r1 = p1.add_run('Bachelor of Networking (Major in Cyber Security)')
r1.bold = True; r1.font.size = Pt(10.5); r1.font.color.rgb = DARK_BLUE

p2 = lc.add_paragraph()
p2.paragraph_format.space_before = Pt(0); p2.paragraph_format.space_after = Pt(4)
r2 = p2.add_run('Melbourne Institute of Technology  ·  Melbourne, VIC 3000')
r2.font.size = Pt(9.5); r2.font.color.rgb = LIGHT_GREY

dp = rc.add_paragraph()
dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
dp.paragraph_format.space_before = Pt(0); dp.paragraph_format.space_after = Pt(0)
dr = dp.add_run('Sep 2023 – Jun 2026\nExpected Graduation')
dr.font.size = Pt(9); dr.italic = True; dr.font.color.rgb = LIGHT_GREY

# ════════════════════════════════════════════════════════════════
# PROJECTS
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Projects')

# ── Featured project ────────────────────────────
ph = doc.add_paragraph()
ph.paragraph_format.space_before = Pt(4)
ph.paragraph_format.space_after  = Pt(1)
pr = ph.add_run('★  AI-Powered Text Recognition & Image Enhancement System')
pr.bold = True; pr.font.size = Pt(11); pr.font.color.rgb = DARK_BLUE
br = ph.add_run('   [AWS + Google Cloud]')
br.bold = True; br.font.size = Pt(9); br.font.color.rgb = MID_BLUE

add_body(doc,
    "Designed and deployed a multi-cloud AI pipeline that automatically detects text in uploaded images, extracts it "
    "via OCR, and — when source images are blurred — applies enhancement algorithms before re-running extraction. "
    "Outputs are saved as separate enhanced image files and structured .txt files, enabling both human review and "
    "downstream automation.",
    size=10, indent=0.3)

ai_bullets = [
    ("AWS Rekognition + Google Cloud Vision API: ", "dual-cloud text detection and OCR with result reconciliation for higher accuracy."),
    ("Serverless event-driven architecture: ",      "AWS Lambda + S3 triggers and Google Cloud Functions + Cloud Storage events for auto-scaling processing."),
    ("Image pre-processing pipeline: ",             "OpenCV & Pillow — blur detection (Laplacian variance), sharpening filters, and contrast enhancement — before OCR on low-quality images."),
    ("Automated end-to-end workflow: ",             "image upload → AI text detection → conditional enhancement → OCR → dual output (enhanced image + .txt) stored to cloud storage."),
    ("Observability: ",                             "CloudWatch + GCP Cloud Logging for monitoring, alerting, and audit trails across both clouds."),
]
for bp, bv in ai_bullets:
    bullet(doc, bv, bold_prefix=bp, size=10)

tools_p = doc.add_paragraph()
tools_p.paragraph_format.space_before = Pt(2)
tools_p.paragraph_format.space_after  = Pt(6)
tr1 = tools_p.add_run('Tools: ')
tr1.bold = True; tr1.font.size = Pt(9.5); tr1.font.color.rgb = DARK_BLUE
tr2 = tools_p.add_run('AWS Lambda, S3, Rekognition, CloudWatch  ·  Google Cloud Vision API, Cloud Functions, Cloud Storage, Cloud Run  ·  Python, OpenCV, Pillow  ·  Git, GitHub Actions')
tr2.font.size = Pt(9.5); tr2.font.color.rgb = MID_BLUE

# ── Other projects ────────────────────────────
other_projects = [
    (
        'Hybrid IT Environment Setup (Azure + VirtualBox)',
        'Microsoft Azure, VirtualBox, Windows Server 2022, Active Directory',
        'Deployed a main server on Microsoft Azure and branch servers in VirtualBox; configured secure AD-based authentication (AD DS, DNS, DHCP) across hybrid environments, simulating a real-world managed-services topology.'
    ),
    (
        'Event-Driven Data Processing Using AWS Serverless',
        'AWS Lambda, S3, EventBridge, CloudWatch · Python',
        'Built a fully serverless real-time data-processing pipeline using event triggers; implemented monitoring dashboards and CloudWatch alarms for proactive alerting.'
    ),
    (
        'Secure Enterprise VPN Network Design & Implementation',
        'Cisco Packet Tracer · IPSec VPN, ACLs',
        'Designed and simulated a site-to-site VPN for branch offices and remote users with encrypted tunnels, strong authentication, and policy-based access control.'
    ),
    (
        'Enterprise Network Design with VLANs & Remote Access',
        'Cisco Packet Tracer · VLAN, VTP, OSPF, WPA2',
        'Built a multi-site enterprise network with VLANs, VTP, wireless access, and secure remote connectivity; applied QoS and inter-VLAN routing.'
    ),
    (
        'Windows Server Administration & User Management',
        'Windows Server 2022, Active Directory, GPMC',
        'Managed Windows Server 2022 with OUs, user accounts, Group Policy Objects (GPOs), and RBAC — directly aligning with Azure AD and ITSM operational practices.'
    ),
    (
        'Digital Forensics Investigation & Evidence Recovery',
        'Autopsy, ProDiscover Basic, FTK Imager',
        'Recovered and analysed deleted data from USB drives; benchmarked tool performance and produced structured forensic reports.'
    ),
    (
        'Bank Management System – Python OOP CLI Application',
        'Python, VS Code, CLI · OOP, file I/O',
        'Developed an OOP-based CLI banking application with account lifecycle management, transaction processing, and persistent storage.'
    ),
]

for title, tools, desc in other_projects:
    ph2 = doc.add_paragraph()
    ph2.paragraph_format.space_before = Pt(6)
    ph2.paragraph_format.space_after  = Pt(1)
    r = ph2.add_run(title)
    r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = DARK_BLUE
    add_body(doc, desc, size=10)
    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(1); tp.paragraph_format.space_after = Pt(4)
    t1 = tp.add_run('Tools: '); t1.bold = True; t1.font.size = Pt(9.5); t1.font.color.rgb = DARK_BLUE
    t2 = tp.add_run(tools);     t2.font.size = Pt(9.5); t2.font.color.rgb = MID_BLUE

# ════════════════════════════════════════════════════════════════
# JOB EXPERIENCE
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Job Experience')

exp_tbl = doc.add_table(rows=1, cols=2)
remove_table_borders(exp_tbl)
exp_tbl.columns[0].width = Inches(4.5)
exp_tbl.columns[1].width = Inches(2.3)
el = exp_tbl.cell(0,0); er = exp_tbl.cell(0,1)
el.paragraphs[0].clear(); er.paragraphs[0].clear()

ep1 = el.add_paragraph()
ep1.paragraph_format.space_before = Pt(0); ep1.paragraph_format.space_after = Pt(1)
ep1r = ep1.add_run('Student Mentor')
ep1r.bold = True; ep1r.font.size = Pt(10.5); ep1r.font.color.rgb = DARK_BLUE

ep2 = el.add_paragraph()
ep2.paragraph_format.space_before = Pt(0); ep2.paragraph_format.space_after = Pt(0)
ep2r = ep2.add_run('Melbourne Institute of Technology  ·  Melbourne, VIC')
ep2r.font.size = Pt(9.5); ep2r.font.color.rgb = LIGHT_GREY

edp = er.add_paragraph()
edp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
edpr = edp.add_run('Jun 2024 – Sep 2025')
edpr.font.size = Pt(9); edpr.font.color.rgb = LIGHT_GREY

exp_bullets = [
    "Acted as a bridge between faculty and students, providing technical guidance on networking and cybersecurity labs — developing customer-focused communication and stakeholder management skills directly applicable to a managed-services role.",
    "Simplified advanced security and networking frameworks for diverse cohorts, demonstrating strong written and verbal communication in technical and non-technical contexts.",
    "Mentored across 10 core units (Networking, Programming, Project Management), enhancing student engagement and subject mastery while managing competing priorities independently.",
]
for b in exp_bullets:
    bullet(doc, b, size=10)

# ════════════════════════════════════════════════════════════════
# CERTIFICATIONS
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Certifications & Training')
certs = [
    'Web Designing with HTML – GitHub',
    'Forage Mastercard Cybersecurity Virtual Experience – Phishing Awareness Campaign Design',
    'CompTIA Security+  (In Progress)',
    'Microsoft Azure Fundamentals (AZ-900)  (Pursuing)',
]
for c in certs:
    bullet(doc, c, size=10)

# ════════════════════════════════════════════════════════════════
# PERSONAL ATTRIBUTES
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Personal Attributes')
attrs = [
    'Strong analytical and problem-solving skills with a security-first mindset',
    'Customer-focused with a professional and adaptable communication style',
    'Excellent teamwork and collaboration in technical and non-technical contexts',
    'Genuine curiosity and eagerness to learn new technologies — particularly Azure',
    'Self-motivated; able to work independently and manage multiple priorities',
    'Proactive approach to identifying, mitigating, and documenting risks',
]
for a in attrs:
    bullet(doc, a, size=10)

# ════════════════════════════════════════════════════════════════
# LANGUAGES
# ════════════════════════════════════════════════════════════════
add_section_heading(doc, 'Languages')
lang_p = doc.add_paragraph()
lang_p.paragraph_format.space_before = Pt(0)
lang_p.paragraph_format.space_after  = Pt(4)
for lang, level in [('English', 'Professional'), ('Hindi', 'Professional'), ('Telugu', 'Native')]:
    lb = lang_p.add_run(lang); lb.bold = True; lb.font.size = Pt(10); lb.font.color.rgb = DARK_BLUE
    lv = lang_p.add_run(f' – {level}     '); lv.font.size = Pt(10); lv.font.color.rgb = BODY_GREY

out_path = '/workspace/Hema_Vardhan_Moganti_Resume.docx'
doc.save(out_path)
print(f"Saved: {out_path}")

"""
BN314 A2 – Questions 5 & 6  (condensed, ≤ 1000 words, preferred style)
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── margins ──────────────────────────────────────────────────
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)

# ── default font ──────────────────────────────────────────────
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

DARK_BLUE  = RGBColor(0x0D, 0x2B, 0x55)
MID_BLUE   = RGBColor(0x1B, 0x4F, 0x8A)
DARK_GREEN = RGBColor(0x14, 0x52, 0x14)

# ── helpers ───────────────────────────────────────────────────
def heading(doc, text, level=1):
    p = doc.add_heading('', level)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.bold = True
    if level == 1:
        r.font.size = Pt(13); r.font.color.rgb = DARK_BLUE
    else:
        r.font.size = Pt(11); r.font.color.rgb = MID_BLUE
    return p

def body(doc, text, space=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.line_spacing = Pt(16)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p

def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = DARK_BLUE
    return p

def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'
    # header row
    hrow = t.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        # blue fill
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1B4F8A')
        tcPr.append(shd)
    # data rows
    for ri, row in enumerate(rows):
        trow = t.rows[ri+1]
        for ci, val in enumerate(row):
            c = trow.cells[ci]
            c.text = val
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(8.5)
    # column widths
    if col_widths:
        for ci, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    return t

def add_image(doc, path, width_in=6.5, cap_text=''):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(path, width=Inches(width_in))
    if cap_text:
        caption(doc, cap_text)

# ═══════════════════════════════════════════════════════════════
#  TITLE PAGE
# ═══════════════════════════════════════════════════════════════
tp = doc.add_heading('', 0)
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run('BN314 – System Architecture\nAssignment 2')
r.font.name = 'Calibri'; r.font.size = Pt(16); r.font.bold = True
r.font.color.rgb = DARK_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run('FreshBite Salads Online Ordering System\nQuestions 5 & 6  |  Trimester T1 2026')
sr.font.size = Pt(11); sr.font.color.rgb = RGBColor(0x44,0x44,0x44)
doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════
#  QUESTION 5 – FIRST-LEVEL DFD
# ═══════════════════════════════════════════════════════════════
heading(doc, 'Question 5 – First-Level Data Flow Diagram (10 Marks)', 1)

body(doc,
    'A Level-1 DFD decomposes the system into major processes, data stores, and data flows '
    'using Yourdon & DeMarco notation. External entities are rectangles, processes are circles, '
    'data stores are open-ended rectangles, and data flows are labelled arrows.')

# 5.1 Processes table
heading(doc, '5.1  Processes', 2)
proc_rows = [
    ('1.0', 'Browse Menu',
     'Retrieves salad details (name, ingredients, nutrition, allergens, price, portions) '
     'from D1 and displays them to the customer.',
     'Browse Request; Salad Info (D1)', 'Salad Details → Customer'),
    ('2.0', 'Manage User Account',
     'Handles registration (with strong-password enforcement), login/logout, email '
     'verification, account status (Active / Suspended / Banned), and reactivation.',
     'Reg/Login Details; User Details (D2); Verification response',
     'Store/Update User (D2); Verification Email; Account Status → Customer'),
    ('3.0', 'Manage Shopping Cart',
     'Allows add/remove/update of salad items per session. Retrieves price from D1; '
     'persists cart to D3.',
     'Add/Remove/Update (Customer); Price Info (D1); Cart Data (D3)',
     'Store Cart Items (D3); Updated Cart → Customer; Cart Items → P4'),
    ('4.0', 'Process Checkout',
     'Verifies authentication (D2), calculates total (items + tax + delivery), collects '
     'delivery address, sends payment to gateway, and passes confirmed order to P5.',
     'Cart Items (P3); Address/Payment (Customer); Auth (D2); Saved Addresses (D5)',
     'Payment Request → Gateway; Store Address (D5); Confirmed Order → P5'),
    ('5.0', 'Manage Orders',
     'Creates order records, tracks status (Pending → Confirmed → Preparing → Out for '
     'Delivery → Delivered), and sends confirmation email.',
     'Confirmed Order (P4); Order Status (D4)',
     'Store Order (D4); Status Updates → Customer; Confirmation Email → Email Service; Order Info → P6'),
    ('6.0', 'Manage Reward Points',
     'Awards 1 point per salad on order completion; checks 15-point eligibility for a '
     'free salad; updates D2.',
     'Completed Order Info (P5); Current Points Balance (D2)',
     'Update Reward Points (D2); Points Balance / Eligibility → Customer'),
]
add_table(doc,
    ['ID', 'Process', 'Description', 'Inputs', 'Outputs'],
    proc_rows,
    col_widths=[0.4, 1.1, 2.0, 1.5, 1.6])

doc.add_paragraph()

# 5.2 Data Stores table
heading(doc, '5.2  Data Stores', 2)
store_rows = [
    ('D1', 'Salad Database',      'Salad ID, name, ingredients, description, allergens, nutrition, price, portion sizes'),
    ('D2', 'User Account DB',     'User ID, username, name, DOB, email, hashed password, account status, failed logins, last activity, reward points'),
    ('D3', 'Shopping Cart',       'Cart ID, session ID, salad references, quantities, portion sizes, item subtotals'),
    ('D4', 'Order Database',      'Order ID, user ref, items, total cost, tax, delivery charge, status, date, payment ref, delivery address'),
    ('D5', 'Address Database',    'Address ID, user ref, street, city, state, postcode, country, address type (Billing/Delivery)'),
]
add_table(doc,
    ['ID', 'Data Store', 'Data Stored'],
    store_rows,
    col_widths=[0.4, 1.3, 4.9])

doc.add_paragraph()

# 5.3 Diagram
heading(doc, '5.3  Level-1 DFD Diagram', 2)
add_image(doc, '/workspace/dfd_v2.png', width_in=6.4,
          cap_text='Figure 1: First-Level Data Flow Diagram – FreshBite Salads Online Ordering System')

# ═══════════════════════════════════════════════════════════════
#  QUESTION 6 – CLASS DIAGRAM
# ═══════════════════════════════════════════════════════════════
heading(doc, 'Question 6 – Class Diagram (15 Marks)', 1)

body(doc,
    'The UML class diagram models the static structure of the system. Each class lists '
    'attributes (name: type) and methods with visibility modifiers (+ public, '
    '- private). Associations are labelled with multiplicity at each end. '
    'CustomSalad inherits from Salad via a generalization (IS-A) relationship.')

# 6.1 Diagram
heading(doc, '6.1  Class Diagram', 2)
add_image(doc, '/workspace/class_v2.png', width_in=6.4,
          cap_text='Figure 2: UML Class Diagram – FreshBite Salads Online Ordering System')

# 6.2 Associations table
heading(doc, '6.2  Associations & Multiplicities', 2)
assoc_rows = [
    ('User', '1', 'ShoppingCart',  '0..*', 'has',          'A user can have zero or many carts (one per session).'),
    ('User', '1', 'Order',         '0..*', 'places',        'A user can place zero or many orders over time.'),
    ('User', '1', 'Address',       '0..*', 'has',           'A user can save zero or many billing/delivery addresses.'),
    ('User', '1', 'RewardPoints',  '1',    'earns',         'Each user has exactly one reward points record.'),
    ('User', '1', 'Verification',  '1',    'requires',      'Each user has exactly one verification record for email activation.'),
    ('ShoppingCart', '1', 'CartItem', '0..*', 'contains',   'A cart holds zero or many cart items.'),
    ('CartItem',     '0..*', 'Salad', '1',    'references', 'Each cart item references exactly one salad.'),
    ('Order', '1', 'CartItem',    '1..*',  'includes',      'An order includes one or more cart items.'),
    ('Order', '1', 'Payment',     '1',     'paid by',       'Each order is settled by exactly one payment record.'),
    ('Order', '1', 'Address',     '1',     'delivered to',  'Each order is linked to one delivery address.'),
]
add_table(doc,
    ['Class A', 'Mult.', 'Class B', 'Mult.', 'Label', 'Description'],
    assoc_rows,
    col_widths=[1.0, 0.45, 1.0, 0.45, 0.85, 2.9])

doc.add_paragraph()

# 6.3 Inheritance
heading(doc, '6.3  Generalisation (Inheritance)', 2)
body(doc,
    'CustomSalad extends Salad using a generalisation (IS-A) relationship, shown as a '
    'hollow-triangle arrow from CustomSalad to Salad. CustomSalad inherits all attributes '
    'and methods of Salad (saladId, name, price, getDetails(), etc.) and adds '
    'customIngredients plus methods addIngredient(), removeIngredient(), and '
    'calculatePrice() to support customer-built salads.')

# ── References ────────────────────────────────────────────────
doc.add_paragraph()
heading(doc, 'References', 1)
refs = [
    'I. Sommerville, Software Engineering, 10th ed. Pearson, 2016.',
    'G. Booch, J. Rumbaugh, and I. Jacobson, The Unified Modeling Language User Guide, 2nd ed. Addison-Wesley, 2005.',
    'E. Yourdon, Modern Structured Analysis. Prentice-Hall, 1989.',
    'Object Management Group, "OMG UML Specification v2.5.1," 2017. [Online]. Available: https://www.omg.org/spec/UML/2.5.1/',
]
for i, ref in enumerate(refs, 1):
    rp = doc.add_paragraph(style='List Number')
    rp.paragraph_format.space_after = Pt(3)
    rr = rp.add_run(f'{ref}')
    rr.font.size = Pt(9.5)

out = '/workspace/BN314_A2_Q5_Q6_v2.docx'
doc.save(out)
print(f'Saved: {out}')

# Word count
words = sum(len(p.text.split()) for p in doc.paragraphs)
print(f'Approximate word count: {words}')

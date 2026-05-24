"""
BN314 A2 – Questions 5 & 6 – Final document
Style mirrors the preferred uploaded document. ≤ 1000 words.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── margins ──────────────────────────────────────────────────
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)

doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

# colour constants
DARK_BLUE  = RGBColor(0x0D, 0x2B, 0x55)
MID_BLUE   = RGBColor(0x1B, 0x4F, 0x8A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
HDR_HEX    = '1B4F8A'

# ── helpers ───────────────────────────────────────────────────
def h1(doc, text):
    p = doc.add_heading('', 1)
    r = p.add_run(text)
    r.font.name = 'Calibri'; r.font.bold = True
    r.font.size = Pt(13); r.font.color.rgb = DARK_BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)

def h2(doc, text):
    p = doc.add_heading('', 2)
    r = p.add_run(text)
    r.font.name = 'Calibri'; r.font.bold = True
    r.font.size = Pt(11); r.font.color.rgb = MID_BLUE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(3)

def body(doc, text, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.line_spacing = Pt(16.5)
    r = p.add_run(text); r.font.size = Pt(11)

def caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    for r in p.runs:
        r.font.size = Pt(9); r.font.italic = True; r.font.color.rgb = MID_BLUE

def add_img(doc, path, w=6.5, cap=''):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(2)
    p.add_run().add_picture(path, width=Inches(w))
    if cap: caption(doc, cap)

def shade_cell(cell, hex_color):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def table(doc, headers, rows, widths=None, hdr_color=HDR_HEX):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = 'Table Grid'
    # header
    hrow = t.rows[0]
    for i, h in enumerate(headers):
        c = hrow.cells[i]; c.text = h
        shade_cell(c, hdr_color)
        for r in c.paragraphs[0].runs:
            r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = WHITE
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
    # data rows
    for ri, row in enumerate(rows):
        trow = t.rows[ri+1]
        fill = 'FFFFFF' if ri%2==0 else 'F2F8FF'
        for ci, val in enumerate(row):
            c = trow.cells[ci]; c.text = val
            shade_cell(c, fill)
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8.5)
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
    if widths:
        for c_idx, w in enumerate(widths):
            for row in t.rows:
                row.cells[c_idx].width = Inches(w)
    return t

# ═══════════════════════════════════════════════════════════════
#  TITLE
# ═══════════════════════════════════════════════════════════════
tp = doc.add_heading('', 0)
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run('BN314 – System Architecture\nAssignment 2')
r.font.name = 'Calibri'; r.font.size = Pt(16); r.font.bold = True
r.font.color.rgb = DARK_BLUE

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run('FreshBite Salads Online Ordering System\n'
                 'Information Management System – Case Study Analysis\n'
                 'Trimester T1 2026  |  Questions 5 & 6')
sr.font.size = Pt(11); sr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════
#  QUESTION 5 – FIRST-LEVEL DFD
# ═══════════════════════════════════════════════════════════════
h1(doc, 'Question 5 – First-Level Data Flow Diagram (10 Marks)')

body(doc,
     'A First-Level Data Flow Diagram (Level-1 DFD) decomposes the single process shown in '
     'the Context Diagram into its major sub-processes. It reveals how data moves between '
     'processes, data stores, and external entities using Yourdon & DeMarco notation: '
     'circles for processes, open-ended rectangles for data stores, rectangles for external '
     'entities, and labelled arrows for data flows.')

# 5.1 External Entities
h2(doc, '5.1  External Entities')
table(doc,
    ['External Entity', 'Description'],
    [
        ('Customer',
         'End user who browses salads, manages their account, adds items to cart, places orders, '
         'and receives status updates. Includes both unregistered (browse only) and registered users.'),
        ('Payment Gateway',
         'Third-party service that processes payments via Credit Card, Debit Card, or PayPal. '
         'Returns payment confirmation or rejection.'),
        ('Email Service',
         'External service that delivers verification emails during registration and order '
         'confirmation emails after successful payment.'),
    ],
    widths=[1.5, 5.1])

# 5.2 Processes
h2(doc, '5.2  Processes')
table(doc,
    ['ID', 'Process', 'Description', 'Key Inputs', 'Key Outputs'],
    [
        ('1.0', 'Browse Menu',
         'Retrieves salad details (name, ingredients, nutrition, allergens, portion sizes, price) '
         'from D1 and presents them to the customer.',
         'Browse Request (Customer); Salad Info (D1)',
         'Salad Details → Customer'),
        ('2.0', 'Manage User Account',
         'Handles registration with strong-password enforcement, login/logout, email verification '
         'via system-generated code, account status management (Active/Suspended/Banned), '
         'and reactivation after suspension.',
         'Reg/Login Details (Customer); User Details (D2); Verification Code Response',
         'Store/Update User (D2); Verification Email; Account Status → Customer'),
        ('3.0', 'Manage Shopping Cart',
         'Manages the session-based cart. Allows customers to add, update quantities, change '
         'portion sizes, and remove items. Retrieves pricing from D1 and persists cart items in D3.',
         'Add/Remove/Update (Customer); Price & Portion Info (D1); Cart Data (D3)',
         'Updated Cart → Customer; Store Cart Items (D3); Cart Items → P4'),
        ('4.0', 'Process Checkout',
         'Verifies user authentication via D2, calculates total (items + tax + delivery), '
         'collects/saves delivery address (D5), sends payment request to Payment Gateway, '
         'and passes confirmed order to P5.',
         'Cart Items (P3); Delivery Address/Payment (Customer); Auth Status (D2); Saved Addresses (D5)',
         'Payment Request → Gateway; Store Address (D5); Confirmed Order → P5'),
        ('5.0', 'Manage Orders',
         'Creates order records in D4 with "Pending" status, updates to "Confirmed" after payment '
         'verification, and tracks status through Preparing → Out for Delivery → Delivered. '
         'Sends confirmation email and forwards completed order info to P6.',
         'Confirmed Payment & Order Details (P4); Order Status (D4)',
         'Store Order (D4); Status Updates → Customer; Confirmation Email → Email Service; Order Info → P6'),
        ('6.0', 'Manage Reward Points',
         'Awards 1 point per salad upon order completion. Checks eligibility for a free salad '
         'at 15 points. Updates the reward points balance in D2.',
         'Completed Order Info (P5); Current Points Balance (D2)',
         'Update Reward Points (D2); Points Balance/Eligibility → Customer'),
    ],
    widths=[0.4, 1.1, 2.1, 1.7, 1.5])

# 5.3 Data Stores
h2(doc, '5.3  Data Stores')
table(doc,
    ['ID', 'Data Store', 'Data Stored'],
    [
        ('D1', 'Salad Database',
         'Salad ID, name, ingredients, description, allergen warnings, nutritional info, price, '
         'portion sizes (Regular, Large, Family)'),
        ('D2', 'User Account Database',
         'User ID, username, first/last name, DOB, email, hashed password, account status '
         '(Active/Suspended/Banned), failed login attempts, last activity date, reward points'),
        ('D3', 'Shopping Cart',
         'Cart ID, session ID, salad references, quantities, portion sizes, item subtotals'),
        ('D4', 'Order Database',
         'Order ID, user ref, order items, total cost, tax, delivery charge, order status, '
         'order date, payment reference, delivery address'),
        ('D5', 'Address Database',
         'Address ID, user ref, street, city, state, postcode, country, address type '
         '(Billing/Delivery)'),
    ],
    widths=[0.4, 1.3, 5.0])

# 5.4 Diagram
h2(doc, '5.4  Level-1 DFD Diagram')
add_img(doc, '/workspace/dfd_final.png', w=6.5,
        cap='Figure 1: First-Level Data Flow Diagram (Level-1 DFD) – '
            'FreshBite Salads Online Ordering System')

# ═══════════════════════════════════════════════════════════════
#  QUESTION 6 – CLASS DIAGRAM
# ═══════════════════════════════════════════════════════════════
h1(doc, 'Question 6 – Class Diagram (15 Marks)')

body(doc,
     'A UML Class Diagram represents the static structure of the system by showing its classes, '
     'their attributes and methods, and the relationships between them. For the FreshBite Salads '
     'Online Ordering System, ten key classes have been identified. Each class includes attributes '
     'with data types and visibility modifiers (- private, + public), methods with return types, '
     'associations with labelled relationships, and multiplicity indicators.')

# 6.1 Class Diagram
h2(doc, '6.1  Class Diagram')
add_img(doc, '/workspace/class_final.png', w=6.5,
        cap='Figure 2: UML Class Diagram – FreshBite Salads Online Ordering System')

# 6.2 Associations
h2(doc, '6.2  Associations & Multiplicities')
table(doc,
    ['#', 'Class A', 'Mult.', 'Class B', 'Mult.', 'Label', 'Description'],
    [
        ('1', 'User', '1', 'ShoppingCart', '0..*', 'has',
         'A user can have zero or many carts (one per session); each cart belongs to exactly one user.'),
        ('2', 'User', '1', 'Order', '0..*', 'places',
         'A registered user can place zero or many orders; each order belongs to exactly one user.'),
        ('3', 'User', '1', 'Address', '0..*', 'has',
         'A user can save zero or many billing/delivery addresses; each address belongs to one user.'),
        ('4', 'User', '1', 'RewardPoints', '1', 'earns',
         'Each registered user has exactly one reward points record tracking accumulated points.'),
        ('5', 'User', '1', 'Verification', '1', 'requires',
         'Each user has exactly one verification record created during registration.'),
        ('6', 'ShoppingCart', '1', 'CartItem', '0..*', 'contains',
         'A shopping cart contains zero or many cart items; each item belongs to exactly one cart.'),
        ('7', 'CartItem', '0..*', 'Salad', '1', 'references',
         'Each cart item references exactly one salad; a salad can appear in many cart items.'),
        ('8', 'Order', '1', 'CartItem', '1..*', 'includes',
         'An order must include at least one cart item linked from the checkout cart.'),
        ('9', 'Order', '1', 'Payment', '1', 'paid by',
         'Each order has exactly one payment transaction; each payment is linked to one order.'),
        ('10', 'Order', '1', 'Address', '1', 'delivered to',
         'Each order is delivered to exactly one delivery address provided at checkout.'),
    ],
    widths=[0.22, 0.9, 0.42, 0.95, 0.42, 0.75, 2.9])

# 6.3 Generalisation
h2(doc, '6.3  Generalisation (Inheritance)')
body(doc,
     'CustomSalad extends Salad using a generalisation (IS-A) relationship, shown in the class '
     'diagram as a hollow-triangle arrow pointing from CustomSalad (child) to Salad (parent). '
     'CustomSalad inherits all attributes and methods from Salad — including saladId, name, '
     'price, and getDetails() — and adds the customIngredients attribute plus '
     'addIngredient(), removeIngredient(), and calculatePrice() methods to support '
     'customer-built salads labeled "Custom Salad" in the system.')

# ── References ────────────────────────────────────────────────
doc.add_paragraph()
h1(doc, 'References')
refs = [
    'I. Sommerville, Software Engineering, 10th ed. Pearson, 2016.',
    'G. Booch, J. Rumbaugh, and I. Jacobson, The Unified Modeling Language User Guide, '
    '2nd ed. Addison-Wesley, 2005.',
    'E. Yourdon, Modern Structured Analysis. Prentice-Hall, 1989.',
    'Object Management Group, "OMG Unified Modeling Language Specification v2.5.1," '
    '2017. [Online]. Available: https://www.omg.org/spec/UML/2.5.1/',
]
for ref in refs:
    rp = doc.add_paragraph(style='List Number')
    rp.paragraph_format.space_after = Pt(3)
    rp.add_run(ref).font.size = Pt(9.5)

out = '/workspace/BN314_A2_Q5_Q6_Final.docx'
doc.save(out)
print(f'Saved: {out}')

from docx import Document as D2
d = D2(out)
w = sum(len(p.text.split()) for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            w += len(cell.text.split())
print(f'Word count (paragraphs + tables): {w}')

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.top_margin    = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin   = Cm(2)
section.right_margin  = Cm(2)

# ── Styles ────────────────────────────────────────────────────────────────────
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)

def set_heading(para, text, level=1, color=RGBColor(0x0e, 0x4d, 0x92)):
    run = para.runs[0] if para.runs else para.add_run(text)
    run.text = text
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 11)
    run.font.color.rgb = color
    para.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    para.paragraph_format.space_after  = Pt(4)

def add_body(doc, text, bold_phrase=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(16)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Calibri'
    return p

def add_bullet(doc, text, indent=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(15)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.name = 'Calibri'
    return p

# ─────────────────────────────────────────────────────────────────────────────
#  Cover heading
# ─────────────────────────────────────────────────────────────────────────────
title = doc.add_heading('', level=0)
title.clear()
run = title.add_run('BN314 – System Architecture | Assignment 2')
run.font.size = Pt(15)
run.font.bold = True
run.font.color.rgb = RGBColor(0x0e, 0x2c, 0x54)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub.add_run('Case Study: FreshBite Salads Online Ordering System')
sub_run.font.size = Pt(12)
sub_run.italic = True
sub_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
sub.paragraph_format.space_after = Pt(14)

# ─────────────────────────────────────────────────────────────────────────────
#  Question 5
# ─────────────────────────────────────────────────────────────────────────────
h5 = doc.add_heading('Question 5: First-Level Data Flow Diagram (DFD)', level=1)
h5.paragraph_format.space_before = Pt(6)
for run in h5.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

add_body(doc,
    "A Level 1 DFD decomposes the central system process of the Context Diagram "
    "into its major functional sub-processes, showing how data moves between those "
    "processes, external entities, and internal data stores. For FreshBite Salads, "
    "seven core processes are identified.")

# Processes table-style bullet list
h5a = doc.add_heading('Major Processes', level=2)
for run in h5a.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

processes = [
    ("P1 – Register & Authenticate User",
     "Handles new-user registration, email verification, login, account suspension "
     "(after 3 failed attempts), and reactivation. Reads/writes DS1 (User Account Store) "
     "and triggers a verification email to the Email Service."),
    ("P2 – Browse & Manage Menu",
     "Retrieves salad details, nutritional information, portion sizes, and prices from "
     "DS2 (Salad/Menu Store). Supports Admin inputs for menu updates."),
    ("P3 – Manage Shopping Cart",
     "Allows both guest and authenticated customers to add, update, or remove cart items. "
     "Persists the session cart in DS3 (Shopping Cart Store)."),
    ("P4 – Checkout & Calculate Total",
     "Loads the cart from DS3, applies taxes and delivery charges, retrieves or saves "
     "the delivery address in DS7, and forwards payment details to P5."),
    ("P5 – Process Payment",
     "Communicates with the external Payment Gateway (credit/debit card or PayPal), "
     "stores the transaction record in DS5, and triggers reward-point accrual via P7."),
    ("P6 – Manage Orders",
     "Creates the initial 'Pending' order in DS4, updates statuses (Confirmed, Preparing, "
     "Out for Delivery, Delivered), and sends a confirmation email to the Email Service."),
    ("P7 – Manage Rewards",
     "Reads and updates reward-point balances in DS6 (1 point per salad purchased; "
     "15 points redeemable for one free salad)."),
]
for title_p, desc in processes:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    run_t = p.add_run(title_p + ': ')
    run_t.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.name = 'Calibri'
    run_d = p.add_run(desc)
    run_d.font.size = Pt(10.5)
    run_d.font.name = 'Calibri'

h5b = doc.add_heading('Data Stores', level=2)
for run in h5b.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

stores = [
    "DS1 – User Account Store: usernames, hashed passwords, statuses, verification codes.",
    "DS2 – Salad/Menu Store: salad IDs, names, descriptions, ingredients, allergens, prices.",
    "DS3 – Shopping Cart Store: session-scoped cart items, quantities, and unit prices.",
    "DS4 – Order Store: order records, statuses, item lines, timestamps.",
    "DS5 – Payment Store: transaction references, amounts, methods, and payment statuses.",
    "DS6 – Reward Points Store: per-user point balances and redemption history.",
    "DS7 – Address Store: saved billing and delivery addresses linked to user accounts.",
]
for s in stores:
    add_bullet(doc, s)

# Insert DFD image
add_body(doc, "The Level 1 DFD below illustrates all processes, data stores, external entities, "
              "and labelled data flows for the FreshBite Salads system.", space_after=4)

img_para = doc.add_paragraph()
img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_para.paragraph_format.space_before = Pt(4)
img_para.paragraph_format.space_after  = Pt(6)
run_img = img_para.add_run()
run_img.add_picture('/workspace/output/dfd_v3.png', width=Inches(6.5))

cap1 = doc.add_paragraph('Figure 1: FreshBite Salads – Level 1 Data Flow Diagram')
cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap1.paragraph_format.space_after = Pt(12)
for run in cap1.runs:
    run.italic = True
    run.font.size = Pt(10)

# ─────────────────────────────────────────────────────────────────────────────
#  Question 6
# ─────────────────────────────────────────────────────────────────────────────
h6 = doc.add_heading('Question 6: Class Diagram', level=1)
for run in h6.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

add_body(doc,
    "The class diagram below captures the static object-oriented structure of the "
    "FreshBite Salads system. Each class is shown with three compartments: name, "
    "attributes (with visibility modifiers), and methods. Relationships carry "
    "multiplicity labels.")

h6a = doc.add_heading('Key Classes, Attributes & Methods', level=2)
for run in h6a.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

classes_desc = [
    ("User «abstract»",
     "Base class for all system users. Holds private fields: userId, username, firstName, "
     "lastName, dateOfBirth, email, and protected fields: password (hashed), accountStatus "
     "(Active/Suspended/Banned), failedLogins, isVerified. "
     "Core methods: register(), login() : boolean, logout(), verifyAccount(), reactivateAccount()."),
    ("RegisteredUser (extends User)",
     "Adds rewardPoints : int and a collection of saved Address objects. "
     "Methods: updateProfile(), manageAddresses(), viewOrderHistory() : List<Order>, "
     "redeemPoints(), saveAddress()."),
    ("GuestUser (extends User)",
     "Holds sessionId and cartId. Limited to browseMenu(), addToCart(), and viewCart(). "
     "Cannot complete checkout until registered."),
    ("Salad",
     "Stores saladId, name, description, ingredients, nutritionalInfo, allergenWarnings, "
     "portionSizes, and a price map keyed on PortionSize (Regular/Large/Family). "
     "isCustom flag distinguishes build-your-own salads. Methods: getPrice(size), "
     "updatePrice(), isAvailable()."),
    ("ShoppingCart",
     "Session-scoped container (cartId, sessionId, userId). Aggregates CartItem objects. "
     "Methods: addItem(), removeItem(), updateQuantity(), calculateTotal(), clearCart()."),
    ("CartItem",
     "Line-item inside a cart: saladId, portionSize, quantity, unitPrice. "
     "Method getSubtotal() returns quantity × unitPrice."),
    ("Order",
     "Created on checkout. Tracks orderId, userId, orderStatus "
     "(Pending → Confirmed → Preparing → OutForDelivery → Delivered), totalAmount, "
     "tax, deliveryCharge, deliveryAddress, and timestamps. "
     "Methods: createOrder(), updateStatus(), calculateTotal(), cancelOrder()."),
    ("OrderItem",
     "Immutable snapshot of a purchased salad line: saladId, portionSize, quantity, "
     "unitPrice. Method getSubtotal()."),
    ("Payment",
     "Holds paymentId, orderId, amount, paymentMethod (CreditCard/DebitCard/PayPal), "
     "paymentStatus (Pending/Verified/Failed), transactionRef, transactionDate. "
     "Methods: processPayment(), verifyPayment(), refundPayment(), getReceipt()."),
    ("Address",
     "Reusable billing/delivery address: street, suburb, city, state, postalCode, "
     "country, addressType. Methods: validate(), toString(), update()."),
    ("RewardProgram",
     "Tracks per-user points (totalPoints, pointsPerSalad=1, pointsThreshold=15). "
     "Methods: addPoints(), deductPoints(), checkEligibility(), redeemFreeSalad(), "
     "getPointsBalance()."),
]

for cls_name, desc in classes_desc:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after  = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    run_t = p.add_run(cls_name + ': ')
    run_t.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.name = 'Calibri'
    run_d = p.add_run(desc)
    run_d.font.size = Pt(10.5)
    run_d.font.name = 'Calibri'

h6b = doc.add_heading('Associations & Multiplicity', level=2)
for run in h6b.runs:
    run.font.color.rgb = RGBColor(0x0e, 0x4d, 0x92)

associations = [
    "User ◁──|> RegisteredUser / GuestUser: Generalisation (inheritance). User is the abstract parent.",
    "RegisteredUser 1 ──▷ 0..* Order (places): One registered user may place zero or many orders.",
    "RegisteredUser 1 ──▷ 0..* Address (has): A user can store multiple billing/delivery addresses.",
    "User / GuestUser 1 ──▷ 1 ShoppingCart (owns): Every session has exactly one cart.",
    "ShoppingCart 1 ──▷ 0..* CartItem (contains): A cart holds zero or more items (composition).",
    "CartItem * ──▷ 1 Salad (references): Many cart-items may reference the same salad.",
    "Order 1 ──▷ 1..* OrderItem (contains): Each order must have at least one line item (composition).",
    "OrderItem * ──▷ 1 Salad (references): Many order-items may reference the same salad.",
    "Order 1 ──▷ 1 Payment (paid via): Every order has exactly one associated payment.",
    "Order 1 ──▷ 1 Address (delivers to): Each order carries one delivery address snapshot.",
    "RegisteredUser 1 ──▷ 1 RewardProgram (enrolled in): Every registered user has one rewards account.",
]
for a in associations:
    add_bullet(doc, a)

# Insert Class Diagram image
add_body(doc, "The complete class diagram with all visibility modifiers, attributes, methods, "
              "associations, and multiplicities is shown below.", space_after=4)

img_para2 = doc.add_paragraph()
img_para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_para2.paragraph_format.space_before = Pt(4)
img_para2.paragraph_format.space_after  = Pt(6)
run_img2 = img_para2.add_run()
run_img2.add_picture('/workspace/output/class_v3.png', width=Inches(6.8))

cap2 = doc.add_paragraph('Figure 2: FreshBite Salads – Class Diagram')
cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap2.paragraph_format.space_after = Pt(8)
for run in cap2.runs:
    run.italic = True
    run.font.size = Pt(10)

# ─── Save ─────────────────────────────────────────────────────────────────────
out_path = '/workspace/output/BN314_A2_Q5_Q6_Answer.docx'
doc.save(out_path)
print(f"Document saved to {out_path}")
